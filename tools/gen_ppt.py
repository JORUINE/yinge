# -*- coding: utf-8 -*-
"""
开题汇报 PPT 生成器（python-pptx）
用途：① 验证 PPT 生成链路可用 ② 产出一份可复用的《音格》开题汇报 PPT
运行：<venv>/python gen_ppt.py <输出pptx路径>
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

W_IN, H_IN = 13.333, 7.5
PURPLE = RGBColor(0x6D, 0x28, 0xD9)
PURPLE_D = RGBColor(0x3B, 0x07, 0x64)
INK = RGBColor(0x1F, 0x24, 0x30)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xF3, 0xEF, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "微软雅黑"


def set_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.makeelement(qn(tag), {"typeface": FONT})
        rPr.append(el)


def add_text(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: [(文本, 字号, 加粗, 颜色), ...]"""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = txt
        set_font(r, size, bold, color)
    return box


def rect(slide, l, t, w, h, fill, line=None, radius=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


prs = Presentation()
prs.slide_width = Inches(W_IN)
prs.slide_height = Inches(H_IN)

# ============ 1. 封面 ============
s = blank(prs)
rect(s, 0, 0, W_IN, H_IN, PURPLE_D)
rect(s, 0, 3.05, W_IN, 0.09, RGBColor(0xA8, 0x55, 0xF7))
add_text(s, 1.0, 1.35, 11.3, 1.7, [
    ("音乐专辑评选与人格测评互动平台设计与实现", 30, True, WHITE),
], align=PP_ALIGN.CENTER)
add_text(s, 1.0, 3.5, 11.3, 2.2, [
    ("—— 毕业论文开题汇报 ——", 16, False, RGBColor(0xC9, 0xB8, 0xF5)),
    ("", 8, False, WHITE),
    ("新媒体学院 · 数字媒体技术 · 2025级（专升本）", 15, False, WHITE),
    ("姓名：胡祖锐　　学号：9920250499　　指导教师：伍春梅", 15, False, WHITE),
    ("2026 年 9 月", 14, False, RGBColor(0xC9, 0xB8, 0xF5)),
], align=PP_ALIGN.CENTER)

SLIDES = [
    ("一、选题背景与研究意义", [
        ("背景其一：比较需求缺乏工具", "用户想系统比较、评选自己喜爱的专辑时，现有音乐应用只提供「推荐—播放」，用户处于被动接受位置"),
        ("背景其二：人格化标签传播力强", "MBTI 类内容持续高传播，说明「结果兼具自我认知与社交表达」的形态有明确产品价值"),
        ("研究意义", "实用：把模糊偏好变成可对决策的擂台与可分享的人格画像　｜　技术：完整前后端分离全栈实践　｜　专业：体现视觉表达与内容生成能力"),
    ]),
    ("二、国内外研究现状与空白", [
        ("国外", "音乐推荐研究强调用户显式反馈的价值；Spotify Wrapped / Apple Music Replay 属「年度回顾型」，用户全程接受结果；音乐对决类社交应用把比较做成了玩法"),
        ("国内", "主流平台以年度听歌报告为主；高校作品集中在「文化数字化展示」与「信息展示类小程序」两类，交互以浏览为主"),
        ("三个空白", "① 参与深度不足：决策行为不被记录、无法累积　② 音乐数据合法性常被忽视，作品难以公开部署　③ 测评结果多为固定文案模板，个性化程度低"),
    ]),
    ("三、研究目标与核心玩法", [
        ("研究目标", "实现一个以「专辑评选」与「音乐人格测评」为核心玩法的互动平台，完成前后端开发、数据库设计、云部署上线"),
        ("玩法一 · 专辑评选", "选歌手/流派/年代 → 自动抓取全部专辑 → 两两对决（可试听 30 秒）→ 小组赛+复活赛+淘汰赛 → 冠军 → 生成「夺冠之路」分享图 → 全球榜单"),
        ("玩法二 · 音乐人格测评", "12 道题（含听感题）→ 人格类型卡片 → 大语言模型生成个性化解读 → 推荐 3 张专辑 → 人格图鉴与占比"),
    ]),
    ("四、技术路线", [
        ("开发方法", "迭代开发，每轮产出可运行版本；引入 AI 辅助编程提效，核心设计与关键逻辑本人完成并验证"),
        ("技术栈", "前端 Vue3 + Vite + Pinia + Element Plus　｜　后端 Node.js + Express　｜　数据库 MongoDB + Mongoose　｜　认证 JWT + 单向加密"),
        ("外部服务与部署", "音乐数据来自 iTunes Search API（免密钥、可公开部署）；文本生成调用大语言模型接口；前端静态托管 + 后端云托管 + 云数据库三端部署"),
    ]),
    ("五、进度安排", [
        ("第 1 周", "需求分析与系统设计：需求分析文档、系统架构图、数据库 ER 图、全部界面设计稿"),
        ("第 2—3 周", "后端与数据库实现 → 前端页面与两条核心玩法实现，完成试听播放与分享图生成"),
        ("第 4—5 周", "大模型接入与后台管理完善 → 三端云部署上线 → 功能/接口/兼容性测试与用户试用反馈"),
        ("第 6 周", "论文撰写与答辩准备：整理图表与过程材料，完成初稿、修改完善、制作演示文稿"),
    ]),
]


def section_slide(title, items):
    s = blank(prs)
    rect(s, 0, 0, W_IN, H_IN, WHITE)
    rect(s, 0, 0, W_IN, 1.15, PURPLE)
    add_text(s, 0.75, 0.22, 11.8, 0.8, [(title, 26, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    y = 1.62
    row_h = min(1.35, (H_IN - 2.1) / max(len(items), 1))
    for label, body in items:
        rect(s, 0.75, y, 11.83, row_h - 0.18, LIGHT)
        rect(s, 0.75, y, 0.075, row_h - 0.18, PURPLE)
        add_text(s, 1.02, y + 0.13, 11.3, row_h - 0.4, [
            (label, 15, True, PURPLE_D),
            (body, 13, False, GREY),
        ])
        y += row_h


for title, items in SLIDES:
    section_slide(title, items)

# ============ 结尾页 ============
s = blank(prs)
rect(s, 0, 0, W_IN, H_IN, PURPLE_D)
add_text(s, 1.0, 3.0, 11.3, 1.5, [
    ("恳请各位老师批评指正", 32, True, WHITE),
    ("谢谢！", 18, False, RGBColor(0xC9, 0xB8, 0xF5)),
], align=PP_ALIGN.CENTER)

out = sys.argv[1] if len(sys.argv) > 1 else "out.pptx"
prs.save(out)
print(f"OK -> {out}")
print(f"幻灯片数: {len(prs.slides.__iter__.__self__._sldIdLst)}")
print(f"尺寸: {W_IN} x {H_IN} 英寸 (16:9)")
